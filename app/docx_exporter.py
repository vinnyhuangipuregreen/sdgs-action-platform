import io
import json
import docx
from app.activity_details import activity_text
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_worksheet_docx(ws: dict) -> io.BytesIO:
    doc = docx.Document()

    # 1. Page Margins (1 inch all sides)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 2. Base Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # Helper: set cell background
    def set_cell_bg(cell, hex_color):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    # Header / Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_org = title_p.add_run("佛光永續學院 × 佛光大學應用科技與設計學院\n")
    run_org.font.size = Pt(11)
    run_org.font.color.rgb = RGBColor(0x05, 0x96, 0x69)
    run_org.font.bold = True

    run_title = title_p.add_run("SDGs 與氣候行動｜連動分析練習單")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x06, 0x4E, 0x3B)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("指導講師：戴慶華 博士 (Dr. ANDERSON) ｜ 核心價值：不只列出編號，說清楚怎麼產生影響")
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    # Metadata Summary Table
    meta_table = doc.add_table(rows=3, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_headers = [
        ("學員姓名", ws.get("user_name", "")),
        ("分會／單位", ws.get("unit_name", "")),
        ("提交時間", str(ws.get("created_at", ""))[:19]),
        ("審核狀態", "已批閱" if ws.get("status") == "reviewed" else "已送出"),
        ("評審成績", f"{ws.get('score')} 分" if ws.get('score') is not None else "待評分"),
        ("中心行動", f"{ws.get('center_action')} ({ws.get('center_custom') or '標準行動'})")
    ]
    
    cells = meta_table.rows[0].cells + meta_table.rows[1].cells + meta_table.rows[2].cells
    for i, (label, val) in enumerate(meta_headers):
        c_lbl = cells[i*2]
        c_val = cells[i*2 + 1]
        c_lbl.text = label
        set_cell_bg(c_lbl, "F0FDF4")
        c_lbl.paragraphs[0].runs[0].font.bold = True
        c_lbl.paragraphs[0].runs[0].font.size = Pt(10)
        c_val.text = str(val)
        c_val.paragraphs[0].runs[0].font.size = Pt(10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 1: Center Action
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(4)
    h1.paragraph_format.keep_with_next = True
    r_h1 = h1.add_run("一、中心行動選擇")
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(0x04, 0x78, 0x57)

    act_names = {"A": "A．一次淨灘", "B": "B．社區二手衣物交換", "C": "C．道場或家裡改用太陽能", "D": "D．自選佛光活動"}
    act_chosen = act_names.get(ws.get("center_action"), ws.get("center_action", "A"))
    p_act = doc.add_paragraph(f"選定行動：{act_chosen}")
    if ws.get("center_custom"):
        p_act.add_run(f" — 自訂名稱：{ws.get('center_custom')}")
    p_act.runs[0].font.bold = True
    p_act.paragraph_format.space_after = Pt(8)
    summary = activity_text(ws.get("activity_details"))
    if summary:
        doc.add_paragraph("活動規劃與成果：").runs[0].font.bold = True
        doc.add_paragraph(summary)

    # Section 2: Radial Analysis
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True
    r_h2 = h2.add_run("二、放射型連動分析（作用機制與可觀察變化）")
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(0x04, 0x78, 0x57)

    radial_data = ws.get("radial_data", [])
    if isinstance(radial_data, str):
        try:
            radial_data = json.loads(radial_data)
        except:
            radial_data = []

    r_table = doc.add_table(rows=1, cols=3)
    r_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    r_headers = ["連動的目標 (SDG)", "怎麼影響？ (作用機制)", "我怎麼看得出來？ (可觀察的變化)"]
    for idx, th in enumerate(r_headers):
        cell = r_table.rows[0].cells[idx]
        cell.text = th
        set_cell_bg(cell, "059669")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.runs[0]
        run.font.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Repeat header row on split
    header_tr = r_table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    for r_idx, item in enumerate(radial_data):
        row = r_table.add_row()
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        c0, c1, c2 = row.cells
        c0.text = item.get("goal", "")
        c1.text = item.get("mechanism", "")
        c2.text = item.get("observable", "")
        if r_idx % 2 == 1:
            set_cell_bg(c0, "F0FDF4")
            set_cell_bg(c1, "F0FDF4")
            set_cell_bg(c2, "F0FDF4")
        for c in (c0, c1, c2):
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 3: Trade-off Analysis
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True
    r_h3 = h3.add_run("三、善意拉扯檢查（負面影響與降低對策）")
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = RGBColor(0xD9, 0x77, 0x06)

    tradeoff_data = ws.get("tradeoff_data", [])
    if isinstance(tradeoff_data, str):
        try:
            tradeoff_data = json.loads(tradeoff_data)
        except:
            tradeoff_data = []

    t_table = doc.add_table(rows=1, cols=3)
    t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_headers = ["可能傷到誰／什麼？", "在什麼情況下會發生？", "我會怎麼降低影響？"]
    for idx, th in enumerate(t_headers):
        cell = t_table.rows[0].cells[idx]
        cell.text = th
        set_cell_bg(cell, "D97706")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.runs[0]
        run.font.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for t_idx, item in enumerate(tradeoff_data):
        row = t_table.add_row()
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        c0, c1, c2 = row.cells
        c0.text = item.get("who", "")
        c1.text = item.get("situation", "")
        c2.text = item.get("how", "")
        if t_idx % 2 == 1:
            set_cell_bg(c0, "FFFBEB")
            set_cell_bg(c1, "FFFBEB")
            set_cell_bg(c2, "FFFBEB")
        for c in (c0, c1, c2):
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 5: Causal Chain
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(12)
    h5.paragraph_format.space_after = Pt(4)
    h5.paragraph_format.keep_with_next = True
    r_h5 = h5.add_run("五、因果型分析")
    r_h5.font.size = Pt(14)
    r_h5.font.bold = True
    r_h5.font.color.rgb = RGBColor(0x43, 0x38, 0xCA)

    causal = ws.get("causal_data", {})
    if isinstance(causal, str):
        try:
            causal = json.loads(causal)
        except:
            causal = {}
            
    causal = causal or {}
    chains = causal.get("chains")
    if chains is None:
        chains = [{**causal, "result_cond": "", "impact": causal.get("result_cond", "")}]
    for chain_number, chain in enumerate(chains or [{}], start=1):
        title = doc.add_paragraph()
        title.paragraph_format.keep_with_next = True
        title.add_run(f"因果鏈 #{chain_number}").font.color.rgb = RGBColor(0xDC, 0x26, 0x26)
        c_table = doc.add_table(rows=6, cols=3)
        c_table.style = "Table Grid"
        c_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        c_rows_data = [
            ("階段", "發生了什麼？", "這一步要成立，需要什麼條件？"),
            ("A 起點", chain.get("start", ""), chain.get("start_cond", "")),
            ("B 中介", chain.get("mid", ""), chain.get("mid_cond", "")),
            ("C 結果", chain.get("result", ""), chain.get("result_cond", "")),
            ("階段", "受到影響是那些人或環境？", "如何衡量此影響？"),
            ("D 影響", chain.get("impact", ""), chain.get("measurement", "")),
        ]
        for idx, values in enumerate(c_rows_data):
            row = c_table.rows[idx]
            for cell, value in zip(row.cells, values):
                cell.text = value
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.space_after = Pt(3)
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                        if idx in (0, 4):
                            run.font.bold = True
            if idx in (0, 4):
                for cell in row.cells:
                    set_cell_bg(cell, "F1F5F9")
            else:
                row.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(0x43, 0x38, 0xCA)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 6: Self-check
    doc.add_heading("六、自我檢核", level=2)
    checks = ws.get("self_checks") or []
    if isinstance(checks, str):
        checks = json.loads(checks)
    check_labels = [
        "我沒有只列 SDGs 編號，每一條線都說得出「怎麼影響」。",
        "我分得清「可能的效益」和「已經證實的成果」。",
        "我找出至少一項負面影響，或可能被遺漏的人。",
        "我提出了做得到改善方法。",
    ]
    for index, label in enumerate(check_labels):
        checked = index < len(checks) and checks[index]
        doc.add_paragraph(f"{'已勾選' if checked else '未勾選'}：{label}")

    # Section 7: Commitment
    h6 = doc.add_paragraph()
    h6.paragraph_format.space_before = Pt(12)
    h6.paragraph_format.space_after = Pt(4)
    h6.paragraph_format.keep_with_next = True
    r_h6 = h6.add_run("七、我的連動承諾宣言")
    r_h6.font.size = Pt(14)
    r_h6.font.bold = True
    r_h6.font.color.rgb = RGBColor(0x04, 0x78, 0x57)

    commit = ws.get("commitment", {})
    if isinstance(commit, str):
        try:
            commit = json.loads(commit)
        except:
            commit = {}

    commit_p = doc.add_paragraph()
    commit_p.paragraph_format.line_spacing = 1.3
    commit_p.add_run("「我願意從 ")
    r_act = commit_p.add_run(commit.get("start_action", "______"))
    r_act.font.bold = True
    r_act.font.underline = True
    commit_p.add_run(" 開始，希望促進 ")
    r_t1 = commit_p.add_run(commit.get("target_1", "______"))
    r_t1.font.bold = True
    r_t1.font.underline = True
    commit_p.add_run(" 與 ")
    r_t2 = commit_p.add_run(commit.get("target_2", "______"))
    r_t2.font.bold = True
    r_t2.font.underline = True
    commit_p.add_run("，也會留意 ")
    r_rk = commit_p.add_run(commit.get("risk_aware", "______"))
    r_rk.font.bold = True
    r_rk.font.underline = True
    commit_p.add_run(" 的風險。一個月後，我會用 ")
    r_mt = commit_p.add_run(commit.get("metric", "______"))
    r_mt.font.bold = True
    r_mt.font.underline = True
    commit_p.add_run(" 來確認自己有沒有做到。」")

    # Section 8: Teacher Review Block
    h7 = doc.add_paragraph()
    h7.paragraph_format.space_before = Pt(16)
    h7.paragraph_format.space_after = Pt(4)
    h7.paragraph_format.keep_with_next = True
    r_h7 = h7.add_run("八、導師審閱成績與指導講評")
    r_h7.font.size = Pt(14)
    r_h7.font.bold = True
    r_h7.font.color.rgb = RGBColor(0x6D, 0x28, 0xD9)

    rev_box = doc.add_table(rows=1, cols=1)
    rev_cell = rev_box.rows[0].cells[0]
    set_cell_bg(rev_cell, "F5F3FF")
    p_rev = rev_cell.paragraphs[0]
    p_rev.paragraph_format.space_before = Pt(6)
    p_rev.paragraph_format.space_after = Pt(6)
    
    score_str = f"{ws.get('score')} 分" if ws.get('score') is not None else "尚未評分"
    comment_str = ws.get("teacher_comment") or "尚無指導講評"
    
    r_s = p_rev.add_run(f"審核評分：{score_str}\n")
    r_s.font.bold = True
    r_s.font.size = Pt(12)
    r_s.font.color.rgb = RGBColor(0x6D, 0x28, 0xD9)
    
    r_c = p_rev.add_run(f"導師回饋：{comment_str}\n\n")
    r_c.font.size = Pt(11)
    
    r_sign = p_rev.add_run("評審單位：佛光永續學院 × 佛光大學應用科技與設計學院 ｜ 講師：戴慶華 博士")
    r_sign.font.size = Pt(9.5)
    r_sign.font.italic = True
    r_sign.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    # Save to buffer
    target_stream = io.BytesIO()
    doc.save(target_stream)
    target_stream.seek(0)
    return target_stream

print("Word Exporter module ready.")
