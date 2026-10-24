"""Student and administrator roster XLSX import and export."""
import io
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

FIELDS = ("name", "email", "email_backup", "phone", "unit_code", "unit_name")
HEADERS = ("學員姓名", "Email（帳號）", "Email（補課平台）", "聯絡電話（密碼）", "分會/分團/童軍團（編號）", "所屬分會/分團/童軍團（完整名稱）")
STAFF_HEADERS = ("管理員姓名", "Email（帳號）", "Email（備用）", "密碼", "單位編號", "單位名稱", "權限角色")


def validate_staff_rows(rows):
    for number, row in enumerate(rows, 1):
        if not all(row.get(field, "").strip() for field in (*FIELDS[:2], "phone", "unit_code", "unit_name", "role")):
            raise HTTPException(400, f"第 {number} 筆管理員資料除 Email（備用）外，各欄位均不可空白")
        if row["role"].strip() not in ("admin", "系統管理員", "teacher", "授課講師/助教"):
            raise HTTPException(400, f"第 {number} 筆的權限角色不正確")
    return rows


def read_roster(content, staff=False):
    headers = STAFF_HEADERS if staff else HEADERS
    fields = (*FIELDS, "role") if staff else FIELDS
    if not content or len(content) > 2 * 1024 * 1024:
        raise HTTPException(400, "請上傳 2 MB 以內的 XLSX 檔案")
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                raise HTTPException(400, "檔案解壓縮後過大，請減少資料筆數")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "無法讀取檔案，請使用有效的 XLSX 格式")
    try:
        rows = workbook.worksheets[0].iter_rows(max_col=len(fields))
        header = next(rows, ())
        if tuple(str(cell.value or "").replace("\ufeff", "").strip() for cell in header) != headers:
            raise HTTPException(400, f"第一張工作表的 A1～{'G1' if staff else 'F1'} 須依序為：" + "、".join(headers))
        students = []
        for number, cells in enumerate(rows, 2):
            if number > 2001:
                raise HTTPException(400, "每次最多匯入 2,000 列資料")
            if all(cell.value is None or str(cell.value).strip() == "" for cell in cells):
                continue
            if any(cell.data_type == "f" for cell in cells):
                raise HTTPException(400, f"第 {number} 列含有公式，請先轉為純文字")
            values = []
            for cell in cells:
                value = cell.value
                if isinstance(value, (int, float)) and not isinstance(value, bool) and value == int(value):
                    value = str(int(value))
                    if cell.number_format and set(cell.number_format) == {"0"}:
                        value = value.zfill(len(cell.number_format))
                values.append(str(value).strip() if value is not None else "")
            student = dict(zip(fields, values))
            if not staff and not all(student[field] for field in ("name", "email", "phone")):
                raise HTTPException(400, f"第 {number} 列的學員姓名、Email（帳號）與聯絡電話不可空白")
            if not staff:
                student["role"] = "student"
            students.append(student)
        if not students:
            raise HTTPException(400, "檔案沒有可匯入的資料")
        if staff:
            validate_staff_rows(students)
        return students
    finally:
        workbook.close()


def export_roster(rows, staff=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "管理員" if staff else "學員名冊"
    sheet.append(STAFF_HEADERS if staff else HEADERS)
    for row in rows:
        values = [str(row.get(field) or "") for field in FIELDS]
        if staff:
            values.append("系統管理員" if row["role"] == "admin" else "授課講師/助教")
        sheet.append(values)
    for row in sheet:
        for cell in row:
            cell.data_type = "s"
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="047857")
    for column, width in zip("ABCDEFG", (18, 34, 34, 22, 30, 48, 24)):
        if column == "G" and not staff:
            break
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
