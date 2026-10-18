"""Database-backed center actions and the two-column XLSX interchange format."""
import io
import json
import unicodedata
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field, ValidationError, model_validator

from app.database import get_db, init_db
from app.activity_details import TEMPLATE_FIELDS


class CategoryItem(BaseModel):
    model_config = {"str_strip_whitespace": True}
    category: str = Field(min_length=1, max_length=80)
    name: str = Field(min_length=1, max_length=300)
    template: dict[str, str] = Field(default_factory=dict)


class CategoryWrite(CategoryItem):
    @model_validator(mode="after")
    def require_templates(self):
        for field, label in TEMPLATE_FIELDS.items():
            if not self.template.get(field, "").strip():
                raise ValueError(f"請填寫：{label}（必填）")
        return self


class TopicItem(BaseModel):
    model_config = {"str_strip_whitespace": True}
    name: str = Field(min_length=1, max_length=80)


def list_topics():
    init_db()
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("""SELECT t.id, t.name, COUNT(a.id) AS activity_count FROM action_topics t
            LEFT JOIN action_catalog a ON a.category = t.name GROUP BY t.id, t.name ORDER BY t.id""")
        return [dict(row) for row in cursor.fetchall()]


def save_topic(item, topic_id=None):
    init_db()
    with get_db() as conn:
        cursor = conn.cursor()
        if topic_id is None:
            cursor.execute("INSERT OR IGNORE INTO action_topics (name) VALUES (?) RETURNING id", (item.name,))
        else:
            cursor.execute("SELECT id FROM action_topics WHERE id = ?", (topic_id,))
            if not cursor.fetchone():
                raise HTTPException(404, "找不到此主題類別，請重新整理")
            cursor.execute("UPDATE OR IGNORE action_topics SET name = ? WHERE id = ? RETURNING id", (item.name, topic_id))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(409, "相同主題類別已存在")
        return {"id": int(row[0]), "name": item.name}


def delete_topic(topic_id):
    init_db()
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("""DELETE FROM action_topics WHERE id = ? AND NOT EXISTS
            (SELECT 1 FROM action_catalog WHERE category = action_topics.name) RETURNING id""", (topic_id,))
        if cursor.fetchone():
            return
        cursor.execute("""SELECT COUNT(a.id) AS activity_count FROM action_topics t
            LEFT JOIN action_catalog a ON a.category = t.name WHERE t.id = ? GROUP BY t.id""", (topic_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(404, "找不到此主題類別，請重新整理")
        raise HTTPException(409, f"該[主題分類]已經有{row['activity_count']}筆對應的[活動內容]，不可刪除！")


def list_actions(search=""):
    init_db()
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT id, category, name, template_data FROM action_catalog ORDER BY id")
        items = [dict(row) for row in cursor.fetchall()]
        for item in items:
            item["template"] = json.loads(item.pop("template_data"))
    terms = unicodedata.normalize("NFKC", search).casefold().split()
    return [item for item in items if all(
        term in unicodedata.normalize("NFKC", f"{item['category']} {item['name']}").casefold()
        for term in terms
    )]


def save_action(item, item_id=None):
    init_db()
    with get_db() as conn:
        cursor = conn.cursor()
        if item_id is None:
            cursor.execute("INSERT OR IGNORE INTO action_catalog (category, name, template_data) VALUES (?, ?, ?) RETURNING id",
                           (item.category, item.name, json.dumps(item.template, ensure_ascii=False)))
        else:
            cursor.execute("SELECT id FROM action_catalog WHERE id = ?", (item_id,))
            if not cursor.fetchone():
                raise HTTPException(404, "找不到此分類項目，請重新整理")
            cursor.execute("UPDATE OR IGNORE action_catalog SET category = ?, name = ?, template_data = COALESCE(?, template_data) WHERE id = ? RETURNING id",
                           (item.category, item.name, json.dumps(item.template, ensure_ascii=False) if 'template' in item.model_fields_set else None, item_id))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(409, "相同主題類別與活動內容已存在")
        return {"id": int(row[0]), **item.model_dump()}


def delete_action(item_id):
    init_db()
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM action_catalog WHERE id = ? RETURNING id", (item_id,))
        if not cursor.fetchone():
            raise HTTPException(404, "找不到此分類項目，請重新整理")


def read_import(content):
    if not content or len(content) > 2 * 1024 * 1024:
        raise HTTPException(400, "請上傳 2 MB 以內的 XLSX 檔案")
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 20 * 1024 * 1024:
                raise HTTPException(400, "檔案解壓縮後過大，請減少資料筆數")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "無法讀取檔案，請使用有效的 XLSX 格式")
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(max_col=2)
        header = next(rows, ())
        if [str(cell.value or "").strip() for cell in header] != ["主題類別", "活動內容"]:
            raise HTTPException(400, "第一張工作表的 A1、B1 須為「主題類別」、「活動內容」")
        items = []
        for number, cells in enumerate(rows, 2):
            if number > 2001:
                raise HTTPException(400, "每次最多匯入 2,000 列資料")
            if all(cell.value is None or str(cell.value).strip() == "" for cell in cells):
                continue
            if any(cell.data_type == "f" for cell in cells):
                raise HTTPException(400, f"第 {number} 列含有公式，請先轉為純文字")
            try:
                item = CategoryItem(category=cells[0].value, name=cells[1].value)
            except ValidationError:
                raise HTTPException(400, f"第 {number} 列資料無效：兩欄皆須為文字且不可空白，類別最多 80 字、活動最多 300 字")
            items.append(item.model_dump())
        if not items:
            raise HTTPException(400, "檔案沒有可匯入的分類項目")
        if "template" in workbook.sheetnames:
            template_rows = workbook["template"].iter_rows(max_col=11)
            header = next(template_rows, ())
            if [str(cell.value or "").strip() for cell in header[:2]] != ["主題類別", "活動內容"]:
                raise HTTPException(400, "template 工作表前兩欄須為「主題類別」、「活動內容」")
            templates = {}
            for number, cells in enumerate(template_rows, 2):
                if number > 2001 or any(cell.data_type == "f" for cell in cells):
                    raise HTTPException(400, "template 工作表不可含公式，且最多 2,000 列")
                if not any(cell.value is not None for cell in cells):
                    continue
                key = tuple(str(cell.value or "").strip() for cell in cells[:2])
                if key in templates:
                    raise HTTPException(400, f"template 第 {number} 列重複，請保留唯一活動範本")
                values = [str(cell.value or "") for cell in cells[2:]]
                templates[key] = dict(zip(TEMPLATE_FIELDS, values)) if any(values) else {}
            for item in items:
                if (item["category"], item["name"]) in templates:
                    item["template"] = templates[(item["category"], item["name"])]
        return items
    finally:
        workbook.close()


def import_actions(items):
    init_db()
    with get_db() as conn:
        cursor = conn.cursor()
        # One SQL statement keeps the import atomic on SQLite and Turso alike.
        cursor.execute("""INSERT INTO action_catalog (category, name, template_data)
            SELECT json_extract(value, '$.category'), json_extract(value, '$.name'), COALESCE(json_extract(value, '$.template'), '{}')
            FROM json_each(?) WHERE true
            ON CONFLICT(category, name) DO UPDATE SET template_data = excluded.template_data
            WHERE excluded.template_data != '{}' AND json(excluded.template_data) != json(action_catalog.template_data)
            RETURNING id""", (json.dumps(items, ensure_ascii=False),))
        imported = len(cursor.fetchall())
    return {"imported_count": imported, "skipped_count": len(items) - imported}


def export_actions():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "分類"
    sheet.append(["主題類別", "活動內容"])
    items = list_actions()
    for item in items:
        sheet.append([item["category"], item["name"]])
    for row in sheet:
        for cell in row:
            cell.data_type = "s"  # Activity text must never become an Excel formula.
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="047857")
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 70
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    template_sheet = workbook.create_sheet("template")
    template_sheet.append(["主題類別", "活動內容", *TEMPLATE_FIELDS.values()])
    for item in items:
        template_sheet.append([item["category"], item["name"], *[item["template"].get(field, "") for field in TEMPLATE_FIELDS]])
    for row in template_sheet:
        for cell in row:
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    template_sheet.freeze_panes = "C2"
    for column in "ABCDEFGHIJK":
        template_sheet.column_dimensions[column].width = 35
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
