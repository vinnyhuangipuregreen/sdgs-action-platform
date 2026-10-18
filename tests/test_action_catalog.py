import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

os.environ["DATABASE_PATH"] = ":memory:"
os.environ.pop("TURSO_DATABASE_URL", None)
os.environ.pop("TURSO_AUTH_TOKEN", None)

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import database
from app.auth import create_access_token
from app.main import app
from app.activity_details import TEMPLATE_FIELDS


def category_item(category, name):
    return {"category": category, "name": name, "template": {field: f"範本：{label}" for field, label in TEMPLATE_FIELDS.items()}}


def xlsx(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    content = io.BytesIO()
    book.save(content)
    return content.getvalue()


class ActionCatalogTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        database.DB_PATH = os.path.join(self.temp.name, "catalog.db")
        database._DB_INITIALIZED = False
        database.init_db()
        with database.get_db() as conn:
            conn.cursor().execute("INSERT INTO admin (id,name,email,phone,password_hash,role) VALUES (1,'Test','test','000','test','admin')")
        with database.get_db() as conn:
            conn.execute("INSERT INTO student (id,name,email,phone,password_hash,role) VALUES (1,'Student','test','000','test','student')")
        self.client = TestClient(app)
        self.admin = self.headers("admin")

    def tearDown(self):
        self.client.close()
        self.temp.cleanup()

    def headers(self, role):
        token = create_access_token({"id": 1, "name": "Test", "email": "test", "unit_name": "Test", "role": role})
        return {"Authorization": "Bearer " + token}

    def upload(self, content, filename="categories.xlsx", headers=None):
        return self.client.post("/api/admin/categories/import", headers=headers or self.admin,
                                files={"file": (filename, content)})

    def test_seeds_once_and_does_not_restore_deleted_items(self):
        self.assertEqual(len(self.client.get("/api/center-actions").json()), 30)
        with database.get_db() as conn:
            conn.cursor().execute("DELETE FROM action_catalog")
        database._DB_INITIALIZED = False
        self.assertEqual(self.client.get("/api/center-actions").json(), [])

    def test_permissions_for_every_admin_operation(self):
        operations = [("GET", "/api/admin/categories", {}),
                      ("POST", "/api/admin/categories", {"json": {"category": "Test", "name": "Test"}}),
                      ("PUT", "/api/admin/categories/1", {"json": {"category": "Test", "name": "Test"}}),
                      ("DELETE", "/api/admin/categories/1", {}),
                      ("GET", "/api/admin/categories/export.xlsx", {}),
                      ("POST", "/api/admin/categories/import", {"files": {"file": ("x.xlsx", xlsx([["主題類別", "活動內容"], ["Test", "Test"]]))}})]
        for method, path, args in operations:
            for headers, expected in [({}, 401), (self.headers("student"), 403)]:
                with self.subTest(method=method, path=path, status=expected):
                    self.assertEqual(self.client.request(method, path, headers=headers, **args).status_code, expected)
        self.assertEqual(self.client.get("/api/admin/categories", headers=self.headers("teacher")).status_code, 200)

    def test_create_update_search_duplicates_and_delete(self):
        item = category_item(" 測試分類 ", " 測試活動 ")
        response = self.client.post("/api/admin/categories", headers=self.admin, json=item)
        self.assertEqual(response.status_code, 200)
        item_id = response.json()["id"]
        self.assertEqual(response.json()["category"], "測試分類")
        self.assertEqual(self.client.post("/api/admin/categories", headers=self.admin, json=item).status_code, 409)
        item = category_item("新主題", "新活動")
        self.assertEqual(self.client.put(f"/api/admin/categories/{item_id}", headers=self.admin, json=item).status_code, 200)
        rows = self.client.get("/api/admin/categories", headers=self.admin, params={"search": "新主題 新活動"}).json()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["id"], item_id)
        self.assertEqual(self.client.delete(f"/api/admin/categories/{item_id}", headers=self.admin).status_code, 200)
        self.assertEqual(self.client.delete(f"/api/admin/categories/{item_id}", headers=self.admin).status_code, 404)
        self.assertEqual(self.client.put(f"/api/admin/categories/{item_id}", headers=self.admin, json=item).status_code, 404)

    def test_validation_and_conflicting_update(self):
        for item in [{"category": " ", "name": "Test"}, {"category": "Test", "name": ""},
                     {"category": "x" * 81, "name": "Test"}, {"category": "Test", "name": "x" * 301}]:
            self.assertEqual(self.client.post("/api/admin/categories", headers=self.admin, json=item).status_code, 422)
        rows = self.client.get("/api/center-actions").json()
        self.assertEqual(self.client.put(f"/api/admin/categories/{rows[0]['id']}", headers=self.admin,
                                        json=category_item(rows[1]["category"], rows[1]["name"])).status_code, 409)

    def test_import_skips_duplicates_and_export_round_trip(self):
        content = xlsx([["主題類別", "活動內容"], ["新類別", "新活動"], ["新類別", "新活動"], [None, None]])
        result = self.upload(content)
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json(), {"imported_count": 1, "skipped_count": 1})
        response = self.client.get("/api/admin/categories/export.xlsx", headers=self.admin)
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.active.freeze_panes, "A2")
        self.assertEqual(book.active.max_row, 32)
        self.assertEqual(list(book.active.values)[-1], ("新類別", "新活動"))
        self.assertEqual(self.upload(response.content).json(), {"imported_count": 0, "skipped_count": 31})

    def test_invalid_import_never_partially_writes(self):
        files = [xlsx([["主題類別", "活動內容"], ["不可部分寫入", "有效列"], ["不完整", None]]),
                 xlsx([["錯誤標題", "活動內容"], ["Test", "Test"]]),
                 xlsx([["主題類別", "活動內容"], ["Test", "=1+1"]]),
                 xlsx([["主題類別", "活動內容"]]), b"invalid", b"x" * (2 * 1024 * 1024 + 1)]
        for content in files:
            self.assertEqual(self.upload(content).status_code, 400)
        self.assertEqual(self.upload(b"invalid", filename="data.csv").status_code, 400)
        self.assertEqual(len(self.client.get("/api/center-actions").json()), 30)

    def test_export_preserves_formula_like_text_as_text(self):
        self.client.post("/api/admin/categories", headers=self.admin, json=category_item("Test", "=1+1"))
        response = self.client.get("/api/admin/categories/export.xlsx", headers=self.admin)
        book = load_workbook(io.BytesIO(response.content), data_only=False)
        self.assertEqual(book.active["B32"].value, "=1+1")
        self.assertEqual(book.active["B32"].data_type, "s")
        self.assertEqual(self.upload(response.content).status_code, 200)

    def test_existing_worksheet_keeps_snapshot(self):
        item = self.client.get("/api/center-actions").json()[0]
        with database.get_db() as conn:
            conn.cursor().execute("""INSERT INTO worksheets
                (user_id, user_name, unit_name, center_action, center_custom, radial_data, tradeoff_data, chat_text, self_checks, commitment)
                VALUES (1, 'Test', 'Test', ?, ?, '[]', '[]', '', '[]', '{}')""", (item["category"], item["name"]))
        response = self.client.put(f"/api/admin/categories/{item['id']}", headers=self.admin, json=category_item("New", "New"))
        self.assertEqual(response.status_code, 200)
        self.client.delete(f"/api/admin/categories/{item['id']}", headers=self.admin)
        with database.get_db() as conn:
            row = conn.cursor().execute("SELECT center_action, center_custom FROM worksheets").fetchone()
        self.assertEqual(tuple(row), (item["category"], item["name"]))

    def test_topics_crud_guard_counts_and_rename_propagation(self):
        topics = self.client.get("/api/admin/topics", headers=self.admin).json()
        actions = self.client.get("/api/center-actions").json()
        topic = topics[0]
        count = sum(row["category"] == topic["name"] for row in actions)
        self.assertEqual(topic["activity_count"], count)
        response = self.client.delete(f"/api/admin/topics/{topic['id']}", headers=self.admin)
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["detail"], f"該[主題分類]已經有{count}筆對應的[活動內容]，不可刪除！")
        self.assertEqual(len(self.client.get("/api/center-actions").json()), len(actions))
        response = self.client.put(f"/api/admin/topics/{topic['id']}", headers=self.admin, json={"name": "重新命名主題"})
        self.assertEqual(response.status_code, 200, response.text)
        renamed = self.client.get("/api/center-actions").json()
        self.assertEqual([r["id"] for r in renamed if r["category"] == "重新命名主題"],
                         [r["id"] for r in actions if r["category"] == topic["name"]])
        self.assertEqual([r["template"] for r in renamed], [r["template"] for r in actions])
        response = self.client.post("/api/admin/topics", headers=self.admin, json={"name": " 空白主題 "})
        self.assertEqual(response.status_code, 200)
        topic_id = response.json()["id"]
        self.assertEqual(response.json()["name"], "空白主題")
        self.assertEqual(self.client.post("/api/admin/topics", headers=self.admin, json={"name": "空白主題"}).status_code, 409)
        self.assertEqual(self.client.put(f"/api/admin/topics/{topic_id}", headers=self.admin, json={"name": "重新命名主題"}).status_code, 409)
        self.assertEqual(self.client.delete(f"/api/admin/topics/{topic_id}", headers=self.admin).status_code, 200)
        database._DB_INITIALIZED = False
        database.init_db()
        self.assertNotIn("空白主題", [r["name"] for r in self.client.get("/api/admin/topics", headers=self.admin).json()])
        self.assertEqual(self.client.delete(f"/api/admin/topics/{topic_id}", headers=self.admin).status_code, 404)
        self.assertEqual(self.client.put(f"/api/admin/topics/{topic_id}", headers=self.admin, json={"name": "不存在"}).status_code, 404)

    def test_import_registers_topics_and_moving_activities_changes_delete_guard(self):
        self.assertEqual(self.upload(xlsx([["主題類別", "活動內容"], ["匯入主題", "匯入活動"]])).status_code, 200)
        topic = next(r for r in self.client.get("/api/admin/topics", headers=self.admin).json() if r["name"] == "匯入主題")
        self.assertEqual(topic["activity_count"], 1)
        action = next(r for r in self.client.get("/api/center-actions").json() if r["category"] == "匯入主題")
        response = self.client.put(f"/api/admin/categories/{action['id']}", headers=self.admin, json=category_item("移入主題", action["name"]))
        self.assertEqual(response.status_code, 200)
        topics = self.client.get("/api/admin/topics", headers=self.admin).json()
        self.assertEqual(next(r for r in topics if r["name"] == "移入主題")["activity_count"], 1)
        self.assertEqual(self.client.delete(f"/api/admin/topics/{topic['id']}", headers=self.admin).status_code, 200)

    def test_each_template_field_is_required_for_create_and_update(self):
        item = category_item("測試", "必填檢核")
        for field in TEMPLATE_FIELDS:
            for invalid in ("", " \n "):
                payload = {**item, "template": {**item["template"], field: invalid}}
                for method, path in (("POST", "/api/admin/categories"), ("PUT", "/api/admin/categories/1")):
                    response = self.client.request(method, path, headers=self.admin, json=payload)
                    self.assertEqual(response.status_code, 422, (method, field))
                    self.assertIn(TEMPLATE_FIELDS[field], response.json()["detail"][0]["msg"])
        self.assertEqual(self.client.post("/api/admin/categories", headers=self.admin, json={"category": "測試", "name": "無範本"}).status_code, 422)

    def test_topic_and_sample_permissions_validation_and_exact_download(self):
        for method, path, args in [("GET", "/api/admin/topics", {}),
                                   ("POST", "/api/admin/topics", {"json": {"name": "Test"}}),
                                   ("PUT", "/api/admin/topics/1", {"json": {"name": "Test"}}),
                                   ("DELETE", "/api/admin/topics/1", {}),
                                   ("GET", "/api/admin/categories/sample.xlsx", {})]:
            for headers, expected in [({}, 401), (self.headers("student"), 403)]:
                self.assertEqual(self.client.request(method, path, headers=headers, **args).status_code, expected)
        for invalid in (" ", "x" * 81):
            self.assertEqual(self.client.post("/api/admin/topics", headers=self.admin, json={"name": invalid}).status_code, 422)
        response = self.client.get("/api/admin/categories/sample.xlsx", headers=self.admin)
        self.assertEqual(response.status_code, 200)
        self.assertIn("category.xlsx", response.headers["content-disposition"])
        self.assertEqual(response.content, (Path(__file__).resolve().parent.parent / "uploadfile/category.xlsx").read_bytes())

    def test_turso_batches_surface_sql_errors(self):
        response = Mock(ok=True)
        response.json.return_value = {"results": [{"type": "error", "error": {"message": "test failure"}}]}
        with patch("app.database.requests.post", return_value=response):
            cursor = database.TursoConnection("https://example.invalid", "test").cursor()
            with self.assertRaisesRegex(Exception, "test failure"):
                cursor.execute_batch(["SELECT 1"])
            with self.assertRaisesRegex(Exception, "test failure"):
                cursor.executemany("SELECT ?", [(1,)])


if __name__ == "__main__":
    unittest.main()
