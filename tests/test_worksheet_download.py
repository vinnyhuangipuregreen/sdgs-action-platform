import io
import unittest
from urllib.parse import unquote
from docx import Document
from tests import test_causal_worksheets as fixtures


class WorksheetDownloadTests(unittest.TestCase):
    setUp = fixtures.CausalWorksheetTests.setUp
    tearDown = fixtures.CausalWorksheetTests.tearDown
    save = fixtures.CausalWorksheetTests.save

    def test_word_filename_uses_unit_code_unit_name_and_author(self):
        worksheet_id, _ = self.save({"chains": self.chains})
        with fixtures.database.get_db() as conn:
            conn.execute("UPDATE admin SET unit_code='A178' WHERE id=1")
            conn.execute("UPDATE worksheets SET user_name=?,unit_name=? WHERE id=?", ("李O立", "松山第五分會", worksheet_id))
        for file_format in ("docx", "doc"):
            response = self.client.get(f"/api/worksheets/{worksheet_id}/export-docx?format={file_format}", headers=self.headers)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(unquote(response.headers["content-disposition"]), f"attachment; filename*=UTF-8''A178_松山第五分會_李O立_連動分析練習單.{file_format}")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"], "application/msword")
        self.assertTrue(response.content.startswith(b"{\\rtf1"))
        self.assertIn(b"\\trowd", response.content)
        for chain in self.chains:
            self.assertIn(chain["measurement"].encode(), response.content)

    def test_filename_replaces_unsupported_characters(self):
        worksheet_id, _ = self.save({"chains": self.chains})
        with fixtures.database.get_db() as conn:
            conn.execute("UPDATE admin SET unit_code='A/178' WHERE id=1")
            conn.execute("UPDATE worksheets SET user_name=?,unit_name=? WHERE id=?", ('李/立', '北投:分會', worksheet_id))
        response = self.client.get(f"/api/worksheets/{worksheet_id}/export-docx", headers=self.headers)
        self.assertEqual(unquote(response.headers["content-disposition"]), "attachment; filename*=UTF-8''A_178_北投_分會_李_立_連動分析練習單.docx")

    def test_docx_export_keeps_its_real_format(self):
        worksheet_id, _ = self.save({"chains": self.chains})
        response = self.client.get(f"/api/worksheets/{worksheet_id}/export-docx", headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"PK"))
        self.assertTrue(Document(io.BytesIO(response.content)).tables)

    def test_required_markers_are_red_in_both_word_formats(self):
        from app.docx_exporter import create_worksheet_docx, create_worksheet_doc
        from docx.shared import RGBColor
        ws = {"center_action": "環保", "activity_details": {
            "name": "海外交流", "start_date": "2026-10-19", "end_date": "2026-10-20",
            "purpose": "目的", "process": "流程", "features": "特色", "audience": "親子",
            "phase": "after", "participants": 0,
        }}
        document = Document(create_worksheet_docx(ws))
        paragraph = next(p for p in document.paragraphs if "開始日期" in p.text)
        for label in ("開始日期", "結束日期", "活動名稱", "1. 活動目的與理念", "2. 活動流程與內容", "3. 活動創意與特色", "4. 預計參與對象"):
            self.assertIn(label + " *：", paragraph.text)
        self.assertIn("實際參與人數：0", paragraph.text)
        markers = [run for run in paragraph.runs if run.text == " *"]
        self.assertEqual(len(markers), 7)
        self.assertTrue(all(run.font.color.rgb == RGBColor(0xDC, 0x26, 0x26) for run in markers))
        rtf = create_worksheet_doc(ws).getvalue()
        self.assertIn(br"\colortbl ;\red220\green38\blue38;", rtf)
        self.assertIn(br"\cf1  *}", rtf)

    def test_doc_download_requires_the_original_authorization(self):
        worksheet_id, _ = self.save({"chains": self.chains})
        response = self.client.get(f"/api/worksheets/{worksheet_id}/export-docx?format=doc")
        self.assertIn(response.status_code, (401, 403))
