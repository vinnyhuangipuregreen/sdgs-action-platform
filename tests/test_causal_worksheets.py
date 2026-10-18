import csv
import io
import json
import os
import tempfile
import unittest

os.environ["DATABASE_PATH"] = ":memory:"
os.environ.pop("TURSO_DATABASE_URL", None)
os.environ.pop("TURSO_AUTH_TOKEN", None)

from docx import Document
from fastapi.testclient import TestClient
from app import database
from app.auth import create_access_token
from app.main import app


class CausalWorksheetTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        database.DB_PATH = os.path.join(self.temp.name, "worksheets.db")
        database._DB_INITIALIZED = False
        database.init_db()
        with database.get_db() as conn:
            conn.cursor().execute("INSERT INTO admin (id,name,email,phone,password_hash,role) VALUES (1,'Test','test@example.invalid','000','test','admin')")
        self.client = TestClient(app)
        self.headers = {"Authorization": "Bearer " + create_access_token({"id": 1, "name": "Test", "email": "test@example.invalid", "role": "admin"})}
        self.chains = [{field: f"{field}-{i}" for field in ("start", "start_cond", "mid", "mid_cond", "result", "result_cond", "impact", "measurement")} for i in range(2)]

    def tearDown(self):
        self.client.close()
        self.temp.cleanup()

    def save(self, causal):
        payload = {"unit_name": "Test", "center_action": "Test", "activity_details": {field:"Test" for field in ("purpose","process","features","audience")}, "radial_data": [], "tradeoff_data": [], "causal_data": causal, "self_checks": [True, False, True, False], "commitment": {}, "status": "draft"}
        response = self.client.post("/api/worksheets", headers=self.headers, json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()["id"], payload

    def test_multiple_chains_round_trip_and_update_without_chat_text(self):
        worksheet_id, payload = self.save({"chains": self.chains})
        url = f"/api/worksheets/{worksheet_id}"
        loaded = self.client.get(url, headers=self.headers).json()
        self.assertEqual(loaded["causal_data"], payload["causal_data"])
        self.assertEqual(loaded["self_checks"], payload["self_checks"])
        self.chains[1]["measurement"] = "updated measurement"
        self.assertEqual(self.client.put(url, headers=self.headers, json=payload).status_code, 200)
        self.assertEqual(self.client.get(url, headers=self.headers).json()["causal_data"]["chains"][1]["measurement"], "updated measurement")
        exported = self.client.get("/api/admin/export/csv", headers=self.headers)
        rows = list(csv.reader(io.StringIO(exported.content.decode("utf-8-sig").lstrip("\ufeff"))))
        self.assertNotIn("聊天室格式文字", rows[0])
        self.assertEqual(json.loads(rows[1][rows[0].index("因果型分析資料")]), payload["causal_data"])

    def test_word_export_includes_both_complete_tables(self):
        worksheet_id, _ = self.save({"chains": self.chains})
        exported = self.client.get(f"/api/worksheets/{worksheet_id}/export-docx", headers=self.headers)
        self.assertEqual(exported.status_code, 200)
        document = Document(io.BytesIO(exported.content))
        tables = [table for table in document.tables if len(table.rows) == 6 and table.cell(0, 0).text == "階段"]
        self.assertEqual(len(tables), 2)
        for i, table in enumerate(tables):
            self.assertEqual(table.cell(5, 0).text, "D 影響")
            self.assertEqual(table.cell(3, 2).text, self.chains[i]["result_cond"])
            self.assertEqual(table.cell(5, 1).text, self.chains[i]["impact"])
            self.assertEqual(table.cell(5, 2).text, self.chains[i]["measurement"])

    def test_legacy_affected_people_move_to_impact_in_word_export(self):
        worksheet_id, _ = self.save({"start": "Legacy start", "result_cond": "Affected people"})
        exported = self.client.get(f"/api/worksheets/{worksheet_id}/export-docx", headers=self.headers)
        document = Document(io.BytesIO(exported.content))
        table = next(table for table in document.tables if len(table.rows) == 6 and table.cell(0, 0).text == "階段")
        self.assertEqual(table.cell(1, 1).text, "Legacy start")
        self.assertEqual(table.cell(3, 2).text, "")
        self.assertEqual(table.cell(5, 1).text, "Affected people")
