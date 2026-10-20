import io
import os
import tempfile
import unittest
from docx import Document
from fastapi.testclient import TestClient
from app import database
from app.auth import create_access_token
from app.main import app


class ActivityDetailsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        database.DB_PATH = os.path.join(self.temp.name, "worksheets.db")
        database._DB_INITIALIZED = False
        database.init_db()
        with database.get_db() as conn:
            for table, role in (("student", "student"), ("admin", "admin")):
                conn.execute(f"INSERT INTO {table} (id,name,email,phone,password_hash,role) VALUES (1,'Test','test@example.invalid','000','test',?)", (role,))
        self.client = TestClient(app)
        self.student = self.headers("student")
        self.admin = self.headers("admin")
        self.payload = {"unit_name":"Test", "center_action":"環保", "center_custom":"淨灘", "activity_details":{"phase":"before","purpose":"目的","process":"流程","features":"特色","audience":"親子"}, "radial_data":[], "tradeoff_data":[], "self_checks":[], "commitment":{}, "status":"draft"}

    def headers(self, role):
        return {"Authorization":"Bearer " + create_access_token({"id":1,"name":"Test","email":"test@example.invalid","role":role})}

    def tearDown(self):
        self.client.close()
        self.temp.cleanup()

    def save(self):
        response = self.client.post("/api/worksheets", headers=self.student, json=self.payload)
        self.assertEqual(response.status_code,200,response.text)
        return response.json()["id"]

    def test_activity_dates_survive_create_update_list_admin_and_exports(self):
        self.payload["activity_details"].update(start_date="2026-10-19", end_date="2026-10-20")
        worksheet_id = self.save()
        url = f"/api/worksheets/{worksheet_id}"
        self.payload["activity_details"]["end_date"] = "2026-10-21"
        self.assertEqual(self.client.put(url, headers=self.student, json=self.payload).status_code, 200)
        for headers in (self.student, self.admin):
            details = self.client.get(url, headers=headers).json()["activity_details"]
            self.assertEqual(details["start_date"], "2026-10-19")
            self.assertEqual(details["end_date"], "2026-10-21")
            listed = self.client.get("/api/worksheets", headers=headers).json()[0]
            self.assertEqual(listed["activity_details"], details)
        exported = self.client.get("/api/admin/export/csv", headers=self.admin).content.decode("utf-8-sig")
        self.assertIn("活動開始日期", exported)
        self.assertIn("2026-10-19", exported)
        self.assertIn("2026-10-21", exported)
        document = Document(io.BytesIO(self.client.get(url + "/export-docx", headers=self.admin).content))
        self.assertIn("開始日期 *：2026-10-19\n結束日期 *：2026-10-21", "\n".join(p.text for p in document.paragraphs))
        word = self.client.get(url + "/export-docx?format=doc", headers=self.admin)
        self.assertIn(b"2026-10-19", word.content)
        self.assertIn(b"2026-10-21", word.content)

    def test_dates_allow_empty_and_same_day_but_reject_incomplete_invalid_or_reversed_ranges(self):
        self.save()
        self.payload["activity_details"].update(start_date="2026-10-19", end_date="2026-10-19")
        self.save()
        for start, end in (("2026-10-20", "2026-10-19"), ("2026-02-30", "2026-03-01"),
                           ("20261019", "2026-10-20"), ("2026-10-19", ""), ("", "2026-10-20")):
            with self.subTest(start=start, end=end):
                self.payload["activity_details"].update(start_date=start, end_date=end)
                self.assertEqual(self.client.post("/api/worksheets", headers=self.student, json=self.payload).status_code, 422)

    def test_retired_fields_are_not_stored_returned_or_exported(self):
        self.payload["activity_description"] = "已刪除的舊說明"
        self.payload["activity_details"]["legacy_description"] = "已刪除的舊說明"
        worksheet_id = self.save()
        url = f"/api/worksheets/{worksheet_id}"
        loaded = self.client.get(url,headers=self.admin).json()
        self.assertNotIn("activity_description",loaded)
        self.assertNotIn("legacy_description",loaded["activity_details"])
        exported = self.client.get("/api/admin/export/csv",headers=self.admin).content.decode("utf-8-sig")
        self.assertNotIn("已刪除的舊說明",exported)
        self.assertNotIn("活動簡介及說明",exported)
        self.assertIn("活動規劃與成果",exported)
        document = Document(io.BytesIO(self.client.get(url+"/export-docx",headers=self.admin).content))
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertNotIn("已刪除的舊說明",text)
        self.assertIn("1. 活動目的與理念 *：目的",text)

    def test_old_description_cannot_bypass_required_planning_fields(self):
        self.payload.pop("activity_details")
        self.payload["activity_description"] = "舊說明"
        self.assertEqual(self.client.post("/api/worksheets",headers=self.student,json=self.payload).status_code,422)

    def test_migration_removes_column_and_nested_legacy_data_without_changing_other_fields(self):
        worksheet_id = self.save()
        with database.get_db() as conn:
            conn.execute("ALTER TABLE worksheets ADD COLUMN activity_description TEXT DEFAULT ''")
            conn.execute("UPDATE worksheets SET activity_description='舊說明', activity_details=json_set(activity_details,'$.legacy_description','舊說明') WHERE id=?", (worksheet_id,))
        for _ in range(2):
            database._DB_INITIALIZED = False
            database.init_db()
        with database.get_db() as conn:
            self.assertNotIn("activity_description",[r["name"] for r in conn.execute("PRAGMA table_info(worksheets)").fetchall()])
        loaded = self.client.get(f"/api/worksheets/{worksheet_id}",headers=self.student).json()
        self.assertEqual(loaded["center_custom"],"淨灘")
        self.assertEqual(loaded["activity_details"]["purpose"],"目的")
        self.assertNotIn("legacy_description",loaded["activity_details"])
