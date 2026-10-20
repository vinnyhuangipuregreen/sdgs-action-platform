import io
import unittest
from datetime import datetime, timedelta, timezone
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from app.accounts import get_accounts_db
from app.auth import verify_password
from app.student_roster import HEADERS
from tests import test_academic_years as fixtures


def workbook_bytes(rows, headers=HEADERS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    result = io.BytesIO()
    workbook.save(result)
    return result.getvalue()


class StudentRosterTests(unittest.TestCase):
    setUp = fixtures.AcademicYearTests.setUp
    tearDown = fixtures.AcademicYearTests.tearDown
    headers = fixtures.AcademicYearTests.headers
    create_year = fixtures.AcademicYearTests.create_year
    switch = fixtures.AcademicYearTests.switch

    def upload(self, content, filename="students.xlsx", headers=None):
        return self.client.post("/api/admin/students/import-xlsx", headers=headers or self.admin,
                                files={"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})

    def test_six_fields_round_trip_preserves_zeroes_and_does_not_duplicate_accounts(self):
        values = ["王小明", "student@example.invalid", "backup@example.invalid", "0920111111", "0012", "北投第一分會"]
        response = self.upload(workbook_bytes([values]))
        self.assertEqual(response.status_code, 200, response.text)
        exported = self.client.get("/api/admin/students/export-xlsx", headers=self.admin)
        self.assertEqual(exported.status_code, 200)
        workbook = load_workbook(io.BytesIO(exported.content))
        self.assertEqual(tuple(cell.value for cell in workbook.active[1]), HEADERS)
        row = next(row for row in workbook.active if row[1].value == values[1])
        self.assertEqual([cell.value for cell in row], values)
        self.assertEqual(row[3].number_format, "@")
        self.assertEqual(row[4].number_format, "@")
        with get_accounts_db() as conn:
            before = conn.execute("SELECT COUNT(*) FROM student").fetchone()[0]
            password = conn.execute("SELECT password_hash FROM student WHERE email=?", (values[1],)).fetchone()[0]
            self.assertTrue(verify_password(values[3], password))
        response = self.upload(exported.content)
        self.assertEqual(response.status_code, 200, response.text)
        with get_accounts_db() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM student").fetchone()[0], before)

    def test_export_filename_uses_taiwan_timestamp_and_is_an_actual_xlsx(self):
        with patch("app.main.datetime") as clock:
            clock.now.return_value = datetime(2026, 10, 20, 23, 4, 5, tzinfo=timezone(timedelta(hours=8)))
            response = self.client.get("/api/admin/students/export-xlsx", headers=self.admin)
            self.assertEqual(clock.now.call_args.args[0].utcoffset(None), timedelta(hours=8))
        self.assertIn("student_20261020230405.xlsx", response.headers["content-disposition"])
        self.assertIn("spreadsheetml.sheet", response.headers["content-type"])
        self.assertTrue(response.content.startswith(b"PK"))

    def test_invalid_uploads_are_rejected_before_any_accounts_are_written(self):
        valid = ["王小明", "student@example.invalid", "", "0920111111", "0012", "北投第一分會"]
        invalid = ["未填電話", "bad@example.invalid", "", "", "", ""]
        cases = [(workbook_bytes([valid]), "students.csv"), (b"not an xlsx", "students.xlsx"),
                 (workbook_bytes([valid], ["wrong"]), "students.xlsx"),
                 (workbook_bytes([valid, invalid]), "students.xlsx"),
                 (workbook_bytes([["=1+1", *valid[1:]]]), "students.xlsx")]
        with get_accounts_db() as conn:
            before = conn.execute("SELECT COUNT(*) FROM student").fetchone()[0]
        for content, filename in cases:
            with self.subTest(filename=filename, size=len(content)):
                self.assertEqual(self.upload(content, filename).status_code, 400)
                with get_accounts_db() as conn:
                    self.assertEqual(conn.execute("SELECT COUNT(*) FROM student").fetchone()[0], before)

    def test_formula_like_names_are_exported_as_text(self):
        values = ["=1+1", "literal@example.invalid", "", "0920111111", "0012", "分會"]
        content = workbook_bytes([])
        workbook = load_workbook(io.BytesIO(content))
        workbook.active.append(values)
        workbook.active['A2'].data_type = 's'
        result = io.BytesIO()
        workbook.save(result)
        self.assertEqual(self.upload(result.getvalue()).status_code, 200)
        exported = self.client.get("/api/admin/students/export-xlsx", headers=self.admin)
        sheet = load_workbook(io.BytesIO(exported.content), data_only=False).active
        cell = next(row[0] for row in sheet if row[1].value == values[1])
        self.assertEqual((cell.value, cell.data_type), (values[0], 's'))

    def test_new_year_import_and_export_do_not_change_the_original_year_roster(self):
        year = self.create_year()
        headers = self.switch(year)
        values = ["新年度學員", "newyear@example.invalid", "", "0912345678", "002", "新年度分會"]
        self.assertEqual(self.upload(workbook_bytes([values]), headers=headers).status_code, 200)
        current = self.client.get("/api/admin/students/export-xlsx", headers=headers)
        old = self.client.get("/api/admin/students/export-xlsx", headers=self.admin)
        current_rows = list(load_workbook(io.BytesIO(current.content)).active.values)
        old_rows = list(load_workbook(io.BytesIO(old.content)).active.values)
        self.assertEqual(current_rows[1:], [tuple(values[:2] + [None] + values[3:])])
        self.assertFalse(any(row[1] == values[1] for row in old_rows[1:]))

    def test_student_accounts_cannot_import_or_export_rosters(self):
        self.assertEqual(self.upload(workbook_bytes([]), headers=self.student).status_code, 403)
        self.assertEqual(self.client.get("/api/admin/students/export-xlsx", headers=self.student).status_code, 403)
